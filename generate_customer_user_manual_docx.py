import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="EFF6FF", border_hex="3B82F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"{title}: ")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_table(tbl, col_widths, header_bg="1E3A8A", alt_bg="F8FAFC"):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(tbl.rows):
        is_header = (i == 0)
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for j, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[j])
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            if is_header:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    for r in p.runs:
                        r.bold = True
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                bg = alt_bg if i % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for r in p.runs:
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

def build_user_manual():
    doc = docx.Document()
    
    # Page setup - Margins 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Captrix Statutory Assurance Platform • Customer User & Implementation Manual")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Strictly Grounded in ICAI Standards on Auditing (SA 200–SA 700) • Companies Act, 2013 Schedule III")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # 1. Main Cover / Header Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    run_brand = title_p.add_run("CAPTRIX ASSURANCE PLATFORM\n")
    run_brand.font.name = "Segoe UI"
    run_brand.font.size = Pt(24)
    run_brand.bold = True
    run_brand.font.color.rgb = RGBColor(0x0F, 0x27, 0x66)

    run_sub = title_p.add_run("Official Customer User Manual & Commercial Implementation Guide")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Document Metadata Pill Table
    tbl_meta = doc.add_table(rows=2, cols=3)
    meta_widths = [2.3, 2.3, 2.4]
    meta_data = [
        [("Product Edition:", " CA Practitioner & Enterprise"), ("Document Version:", " 4.0 (Production)"), ("Statutory Alignment:", " ICAI SAs & Companies Act")],
        [("Supported ERPs:", " Tally, Zoho, SAP, Busy, Marg"), ("Target Users:", " CAs, Audit Managers, CFOs"), ("Processing Engine:", " Deterministic Zero-Error")]
    ]
    for r_idx, row in enumerate(meta_data):
        for c_idx, (k, v) in enumerate(row):
            cell = tbl_meta.cell(r_idx, c_idx)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r1 = p.add_run(k)
            r1.bold = True
            r1.font.name = "Segoe UI"
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            r2 = p.add_run(v)
            r2.font.name = "Segoe UI"
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    for i in range(2):
        for j in range(3):
            tbl_meta.cell(i, j).width = Inches(meta_widths[j])

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 2. Executive Value Proposition (Why Buy Captrix?)
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Executive Summary & Commercial Value Proposition")
    r_h1.font.name = "Segoe UI"
    r_h1.font.size = Pt(14)
    r_h1.bold = True
    r_h1.font.color.rgb = RGBColor(0x0F, 0x27, 0x66)

    p_intro = doc.add_paragraph(
        "Captrix is India's leading automated statutory assurance and audit intelligence software designed "
        "specifically for practicing Chartered Accountant firms, audit partners, and corporate finance controllers. "
        "Unlike generic spreadsheet templates or black-box systems, Captrix executes audits deterministically, "
        "adhering 100% to ICAI Standards on Auditing (SA 200–SA 700) and Schedule III of the Companies Act, 2013."
    )
    p_intro.paragraph_format.space_after = Pt(8)

    add_callout(
        doc,
        "THE ZERO-ERROR MATHEMATICAL GUARANTEE",
        "If the uploaded accounting records are mathematically valid, Captrix guarantees 100% accurate statutory calculations "
        "with an exact Trial Balance difference of ₹0.00 and Balance Sheet equilibrium (Total Assets = Total Equity & Liabilities). "
        "No artificial balancing figures, no dummy entries, and no hallucinations.",
        bg_hex="ECFDF5",
        border_hex="10B981"
    )

    tbl_roi = doc.add_table(rows=5, cols=3)
    roi_widths = [2.0, 2.7, 2.3]
    roi_data = [
        ["Audit Workflow Stage", "Traditional Manual Audit (Days/Weeks)", "Captrix Autonomous Engine"],
        ["Schedule III Mapping", "7 to 10 Days of manual ledger linking", "Instant (Under 60 seconds)"],
        ["Bank Reconciliation (BRS)", "3 to 5 Days verifying physical bank statements", "Automated 4-Way Matching in seconds"],
        ["MSME Sec 43B(h) Audit", "High risk of overlooking 45-day vendor invoices", "Automated 45-day subledger aging bucket"],
        ["Working Paper Compilation", "Scattered Excel files prone to peer review flags", "Cryptographic SA 230 Working Paper Vault"]
    ]
    for r_idx, row in enumerate(roi_data):
        for c_idx, val in enumerate(row):
            tbl_roi.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(tbl_roi, roi_widths, header_bg="1E3A8A", alt_bg="F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3. User Personas & Role Matrix
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. User Personas & Role-Based Responsibilities (Who Uses What?)")
    r_h2.font.name = "Segoe UI"
    r_h2.font.size = Pt(14)
    r_h2.bold = True
    r_h2.font.color.rgb = RGBColor(0x0F, 0x27, 0x66)

    doc.add_paragraph(
        "To comply with ICAI SA 220 (Quality Control for an Audit of Financial Statements), Captrix enforces "
        "a strict separation of duties between preparatory staff and reviewing engagement partners:"
    )

    tbl_roles = doc.add_table(rows=5, cols=4)
    role_widths = [1.5, 1.8, 2.3, 1.4]
    role_data = [
        ["User Persona", "Primary Objective", "Core Responsibilities in Captrix", "Accessible Screens"],
        ["Engagement Partner (FCA / CA)", "Statutory Opinion & Legal Sign-off", "Affixes official ICAI UDIN, approves modified opinions, signs Final Audit Dossier", "Final Audit File, Review Queue, Command Center"],
        ["Audit Manager", "Quality Control & Scoping", "Sets SA 320 materiality benchmarks, reviews assistant workpapers, signs off BRS matches", "Auditor Suite, Reconciliation, Review Queue"],
        ["Audit Senior / Article", "Substantive Field Testing", "Executes 6-pillar audit matrix, investigates Sec 40A(3) cash entries, drafts Adjusting Journal Entries", "Upload, Auditor Suite, Findings & Exceptions"],
        ["Client Finance / Operator", "Data Staging & Ingestion", "Exports accounts from Tally/Zoho/SAP, stages files via drag-and-drop, monitors pipeline", "Upload & Analyze, Processing Status"]
    ]
    for r_idx, row in enumerate(role_data):
        for c_idx, val in enumerate(row):
            tbl_roles.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(tbl_roles, role_widths, header_bg="0F2766", alt_bg="F1F5F9")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. Input Preparation Guide
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. Input Preparation Guide (How to Feed Client Data)")
    r_h3.font.name = "Segoe UI"
    r_h3.font.size = Pt(14)
    r_h3.bold = True
    r_h3.font.color.rgb = RGBColor(0x0F, 0x27, 0x66)

    doc.add_paragraph(
        "Captrix accepts either a Single Master Excel Workbook (.xlsx) with multiple sheets, individual CSV exports, "
        "or a compressed .ZIP file. The 7 recognized client account sheets are:"
    )

    tbl_inputs = doc.add_table(rows=8, cols=4)
    input_widths = [1.6, 1.3, 2.6, 1.5]
    input_data = [
        ["Sheet / File Name", "Mandatory?", "Recognized Column Aliases (Auto-Normalized)", "Audit Impact"],
        ["Trial Balance", "MANDATORY", "Particulars, Account Name, Ledger, Debit, Credit, Closing Balance", "Schedule III Balance Sheet & P&L"],
        ["General Ledger", "MANDATORY", "Date, Voucher No, Vch Type, Account, Debit, Credit, Narration", "Sec 40A(3), Benford's Law, Scrutiny"],
        ["Bank Statement", "MANDATORY", "Txn Date, Value Date, Narration, Withdrawal (Dr), Deposit (Cr), Balance", "BRS 4-Way Reconciliation"],
        ["GSTR-2B Statement", "RECOMMENDED", "GSTIN of Supplier, Trade Name, Invoice No, Taxable Value, CGST, SGST, IGST", "GST Rule 36(4) ITC Reconciliation"],
        ["Purchase Register", "RECOMMENDED", "Invoice No, Supplier Name, Invoice Date, Taxable Value, GSTIN", "Inward Supply vs GSTR-2B Match"],
        ["Fixed Asset Register", "RECOMMENDED", "Asset Description, Date of Acquisition, Original Cost, Depreciation Rate", "CARO 2020 Clause 3(i) FAR Audit"],
        ["Subledger Aging", "RECOMMENDED", "Party Name, MSME Status (Micro/Small), Bill Date, Due Date, Outstanding Amount", "Section 43B(h) 45-Day Overdue Check"]
    ]
    for r_idx, row in enumerate(input_data):
        for c_idx, val in enumerate(row):
            tbl_inputs.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(tbl_inputs, input_widths, header_bg="1E3A8A", alt_bg="F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    add_callout(
        doc,
        "DYNAMIC ADAPTIVE AUDIT MODE",
        "If a client does not maintain Fixed Assets or GSTR-2B data, Captrix automatically suppresses non-applicable procedures "
        "without error. The system dynamically scales to focus strictly on the records provided.",
        bg_hex="EFF6FF",
        border_hex="3B82F6"
    )

    # 5. Universal Account Classification (Small MSME to Large IT)
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. Universal Account Classification (MSME to Large IT Enterprise)")
    r_h4.font.name = "Segoe UI"
    r_h4.font.size = Pt(14)
    r_h4.bold = True
    r_h4.font.color.rgb = RGBColor(0x0F, 0x27, 0x66)

    doc.add_paragraph(
        "Captrix is pre-calibrated to classify 99+ real-world account types with 100% deterministic accuracy:"
    )

    tbl_class = doc.add_table(rows=6, cols=3)
    class_widths = [2.0, 2.8, 2.2]
    class_data = [
        ["Business Scale Tier", "Sample Ledger Account Heads Tested", "Schedule III Classification"],
        ["Small MSME & Proprietorship", "Coolie, Hamali, Cartage, Loading/Unloading Charges", "Cost of Materials Consumed (Direct Costs)"],
        ["Local Office Kutty Accounts", "Office Tea/Snacks, Pooja, Shop Rent, Xerox, Water Cans", "Other Expenses / Staff Welfare Expenses"],
        ["Employee Payroll & Benefits", "Basic, HRA, DA, Special Allowance, PF, ESIC, Gratuity", "Employee Benefit Expenses & Current Liabilities"],
        ["IT & Tech Enterprises", "AWS Cloud, Azure, GCP Compute, MongoDB, OpenAI API", "Other Expenses - Cloud Hosting & Technology"],
        ["Developer Tools & Hardware", "GitHub, Jira, Slack, Figma, MacBook Pros, Dell Servers", "SaaS Subscriptions (Exp) & PPE (Assets)"]
    ]
    for r_idx, row in enumerate(class_data):
        for c_idx, val in enumerate(row):
            tbl_class.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(tbl_class, class_widths, header_bg="0F2766", alt_bg="F1F5F9")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 6. Step-by-Step User Walkthrough
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(6)
    r_h5 = h5.add_run("5. Step-by-Step Operational Walkthrough (The 4 Steps)")
    r_h5.font.name = "Segoe UI"
    r_h5.font.size = Pt(14)
    r_h5.bold = True
    r_h5.font.color.rgb = RGBColor(0x0F, 0x27, 0x66)

    doc.add_paragraph(
        "Auditors complete the entire statutory audit lifecycle through 4 straightforward steps:"
    )

    steps = [
        ("Step 1: Staging the Client Workbook (/upload)",
         "Select the client company name and financial year (FY 2024-25). Drag and drop the master Excel workbook. "
         "Captrix computes the SHA-256 cryptographic hash and checks non-zero byte integrity."),
        ("Step 2: Monitoring the 14-Stage Pipeline (/processing)",
         "Captrix executes 14 automated stages in real time: Hash verification, column normalization, cell extraction, "
         "Schedule III calculations, BRS matching, statutory testing (Sec 40A(3), Sec 43B(h)), and report packaging."),
        ("Step 3: CA Auditor Suite Scrutiny (/auditor-suite & /exceptions)",
         "Inspect the 6-Pillar Statutory Matrix. Review and resolve flagged exceptions with auditor notes. "
         "Approve BRS timing differences and verified subledger balances."),
        ("Step 4: Quality Review & Sign-Off (/final-audit-file)",
         "Review the 14-section Consolidated Audit Dossier. Enter CA Membership Number, FRN, and official 18-digit ICAI UDIN. "
         "Instantly download the formal PDF Audit Report, 8-Tab Excel schedules, and complete .ZIP audit package.")
    ]

    for title, desc in steps:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(4)
        sp.paragraph_format.space_after = Pt(4)
        r_st = sp.add_run(f"• {title}: ")
        r_st.bold = True
        r_st.font.name = "Segoe UI"
        r_st.font.size = Pt(10)
        r_st.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
        r_sd = sp.add_run(desc)
        r_sd.font.name = "Segoe UI"
        r_sd.font.size = Pt(9.5)
        r_sd.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 7. Deliverables & Exports
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(16)
    h6.paragraph_format.space_after = Pt(6)
    r_h6 = h6.add_run("6. Audit Deliverables & Export Formats")
    r_h6.font.name = "Segoe UI"
    r_h6.font.size = Pt(14)
    r_h6.bold = True
    r_h6.font.color.rgb = RGBColor(0x0F, 0x27, 0x66)

    tbl_exp = doc.add_table(rows=5, cols=3)
    exp_widths = [2.2, 2.0, 2.8]
    exp_data = [
        ["Deliverable Package", "Format", "Commercial Purpose & Contents"],
        ["Consolidated Audit Report", "Formal PDF (ReportLab)", "Complete multi-page statutory audit report with ICAI watermark, running page numbers, and CA signature."],
        ["Consolidated Audit Schedules", "Excel (.xlsx, 8 Tabs)", "8 distinct tabs: Balance Sheet, P&L, 11 Ratios, Trial Balance, BRS, FAR Schedule, MSME Aging, and Exceptions."],
        ["Standardized Audit Data", "CSV (.csv)", "Machine-readable raw data for regulatory filing and client ERP re-import."],
        ["Cryptographic Audit Archive", "ZIP Package (.zip)", "Complete electronic working paper folder (SA 230) containing all original evidence, hashes, and signed outputs."]
    ]
    for r_idx, row in enumerate(exp_data):
        for c_idx, val in enumerate(row):
            tbl_exp.cell(r_idx, c_idx).paragraphs[0].text = val
    style_table(tbl_exp, exp_widths, header_bg="1E3A8A", alt_bg="F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Closing Callout
    add_callout(
        doc,
        "CUSTOMER SUPPORT & SLA",
        "Captrix provides comprehensive onboarding, standard chart of accounts templates, and 24/7 technical assurance support. "
        "Engineered for practicing Chartered Accountants and Enterprise Finance teams across India.",
        bg_hex="F1F5F9",
        border_hex="475569"
    )

    out_path = os.path.join(r"d:\Captrix Ai", "Captrix_Official_Customer_User_Manual_and_Sales_Guide.docx")
    doc.save(out_path)
    print(f"Successfully generated DOCX manual at: {out_path}")

if __name__ == "__main__":
    build_user_manual()
