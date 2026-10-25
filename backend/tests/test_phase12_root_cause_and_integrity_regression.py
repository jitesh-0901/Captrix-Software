import pytest
import io
import openpyxl
from decimal import Decimal
from datetime import date
from django.utils import timezone
from django.contrib.auth.models import User
from rest_framework.test import APIClient

from apps.companies.models import Company, FinancialYear, Engagement, EngagementType, EntityType, EngagementStatus
from apps.accounting.models import ScheduleIIIGroup, TrialBalanceSnapshot, LedgerAccount
from apps.accounting.services import TrialBalanceService
from apps.audit.models import (
    SubledgerAgingItem, ExceptionItem, ExceptionSeverity, ExceptionStatus,
    ProcedureExecution, ProcedureStatus
)
from apps.reconciliation.models import ReconciliationRun, ReconciliationStatus, ReconciliationType
from apps.ingestion.models import Document, DocumentType, DocumentStatus
from apps.reporting.models import Workpaper, WorkpaperStatus, ProposedAdjustment, AdjustmentStatus
from apps.reporting.report_builder import ConsolidatedAuditReportBuilder
from apps.reporting.exporters import ReportExporterService
from apps.reporting.integrity_gate import ReportIntegrityGateService
from apps.audit.finalization_service import AuditFinalizationService
from apps.companies.completion_service import EngagementCompletionService


@pytest.mark.django_db
class TestPhase12RootCauseAndIntegrityRegression:
    """
    Phase 12 Comprehensive Root-Cause and Regression Test Suite.
    Enforces non-negotiable accounting integrity, reconciliation gates,
    statutory audit opinion constraints, and report status consistency.
    """

    @pytest.fixture
    def setup_engagement(self):
        company = Company.objects.create(
            name="Vanguard Technologies Private Limited",
            cin="U72200MH2022PTC654321",
            pan="AABCV1234F",
            entity_type=EntityType.PVT_LTD
        )
        fy = FinancialYear.objects.create(
            company=company,
            label="FY 2025-26",
            start_date=date(2025, 4, 1),
            end_date=date(2026, 3, 31)
        )
        partner = User.objects.create_user(
            username="senior_audit_partner_ca",
            password="securepassword123",
            is_staff=True
        )
        engagement = Engagement.objects.create(
            company=company,
            financial_year=fy,
            title="Statutory Audit FY 2025-26",
            engagement_type=EngagementType.STATUTORY,
            overall_materiality=Decimal("500000.00"),
            performance_materiality=Decimal("375000.00"),
            lead_partner=partner,
            is_finalized=False
        )

        # Baseline balanced Trial Balance
        tb_rows = [
            {'account_code': 'REV101', 'account_name': 'Sales Revenue', 'closing_credit_inr': '12000000.00', 'schedule_iii_group': ScheduleIIIGroup.REVENUE_OPERATIONS},
            {'account_code': 'EXP201', 'account_name': 'Cost of Materials Consumed', 'closing_debit_inr': '7000000.00', 'schedule_iii_group': ScheduleIIIGroup.COST_OF_MATERIALS},
            {'account_code': 'EXP202', 'account_name': 'Employee Benefit Expenses', 'closing_debit_inr': '2000000.00', 'schedule_iii_group': ScheduleIIIGroup.EMPLOYEE_BENEFITS},
            {'account_code': 'EQ301', 'account_name': 'Equity Share Capital', 'closing_credit_inr': '3000000.00', 'schedule_iii_group': ScheduleIIIGroup.SHARE_CAPITAL},
            {'account_code': 'CL302', 'account_name': 'Trade Payables GL Control', 'closing_credit_inr': '2500000.00', 'schedule_iii_group': ScheduleIIIGroup.TRADE_PAYABLES},
            {'account_code': 'CA401', 'account_name': 'Trade Receivables GL Control', 'closing_debit_inr': '3500000.00', 'schedule_iii_group': ScheduleIIIGroup.TRADE_RECEIVABLES},
            {'account_code': 'CA402', 'account_name': 'Bank Account Current', 'closing_debit_inr': '5000000.00', 'schedule_iii_group': ScheduleIIIGroup.CASH_EQUIVALENTS},
        ]
        # Total Credits: 12M + 3M + 2.5M = 17.5M. Total Debits: 7M + 2M + 3.5M + 5M = 17.5M.
        TrialBalanceService.calculate_trial_balance(tb_rows, as_of_date=fy.end_date, engagement=engagement)

        return {
            'company': company,
            'fy': fy,
            'partner': partner,
            'engagement': engagement
        }

    # -------------------------------------------------------------------------
    # TEST 1: Trial Balance Debit/Credit Totals Equilibrium
    # -------------------------------------------------------------------------
    def test_trial_balance_debit_credit_totals_equilibrium(self, setup_engagement):
        engagement = setup_engagement['engagement']
        tb = engagement.trial_balances.first()
        assert tb is not None
        assert tb.is_balanced is True
        assert tb.difference_amount == Decimal("0.00")
        assert tb.total_debits == Decimal("17500000.00")
        assert tb.total_credits == Decimal("17500000.00")

        # Now induce a deliberate double-entry variance
        unbalanced_rows = [
            {'account_code': 'REV101', 'account_name': 'Sales Revenue', 'closing_credit_inr': '10000000.00'},
            {'account_code': 'EXP201', 'account_name': 'Operating Expenses', 'closing_debit_inr': '9000000.00'},
        ]
        unbalanced_tb = TrialBalanceService.calculate_trial_balance(
            unbalanced_rows, as_of_date=setup_engagement['fy'].end_date
        )
        assert unbalanced_tb['is_balanced'] is False
        assert abs(Decimal(str(unbalanced_tb['difference_amount']))) == Decimal("1000000.00")

    # -------------------------------------------------------------------------
    # TEST 2: Balance Sheet Equation ($Assets = Liabilities + Equity$)
    # -------------------------------------------------------------------------
    def test_balance_sheet_equation_reconciliation(self, setup_engagement):
        engagement = setup_engagement['engagement']
        tb = engagement.trial_balances.first()
        fs = TrialBalanceService.compile_financial_statements(tb.accounts_snapshot)

        bs = fs['balance_sheet']
        tot_assets = Decimal(str(bs['assets']['total_assets']))
        tot_eq_liab = Decimal(str(bs['equity_and_liabilities']['total_equity_and_liabilities']))
        variance = abs(tot_assets - tot_eq_liab)

        assert variance < Decimal("1.00"), f"Balance Sheet disequilibrium: Assets ₹{tot_assets} != Equity+Liabilities ₹{tot_eq_liab}"

    # -------------------------------------------------------------------------
    # TEST 3: Unresolved Bank Reconciliation Blocks Finalization and Opinion
    # -------------------------------------------------------------------------
    def test_unresolved_bank_reconciliation_blocks_finalization_and_opinion(self, setup_engagement):
        engagement = setup_engagement['engagement']
        partner = setup_engagement['partner']

        # Create an UNRESOLVED bank reconciliation run
        brs = ReconciliationRun.objects.create(
            engagement=engagement,
            recon_type=ReconciliationType.BANK,
            status=ReconciliationStatus.VARIANCE_DETECTED,
            as_of_date=setup_engagement['fy'].end_date,
            matched_count=45,
            unmatched_count=4,  # 4 unmatched bank items
            unreconciled_difference=Decimal("125000.00"),
            summary_data={'stale_cheques_count': 1}
        )

        # 1. AuditFinalizationService must block finalization
        fin_readiness = AuditFinalizationService.evaluate_finalization_readiness(engagement, user=partner)
        assert fin_readiness['can_finalize'] is False
        assert any("Bank reconciliation unresolved" in b for b in fin_readiness['blockers'])

        # 2. Permitted exceptions must NOT be allowed to override core BRS gate
        engagement.permitted_exceptions = [{'code': 'BANK', 'reason': 'Auditor note trying to bypass'}]
        engagement.save()
        fin_readiness_after = AuditFinalizationService.evaluate_finalization_readiness(engagement, user=partner)
        assert fin_readiness_after['can_finalize'] is False
        assert any("Bank reconciliation unresolved" in b for b in fin_readiness_after['blockers'])

        # 3. Attempting to finalize via AuditFinalizationService must fail closed
        fin_result = AuditFinalizationService.finalize_engagement(engagement, user=partner)
        assert fin_result['success'] is False
        engagement.refresh_from_db()
        assert engagement.is_finalized is False

        # 4. ReportIntegrityGateService must disallow clean statutory opinion
        integrity_res = ReportIntegrityGateService.evaluate_integrity(engagement)
        assert integrity_res['is_ready_for_final_issue'] is False
        assert integrity_res['opinion_allowed'] is False

        # 5. Report builder and exports must truthfully declare BRS unresolved
        rep_data = ConsolidatedAuditReportBuilder.build_report(engagement)
        assert any("Bank reconciliation: UNRESOLVED" in item for item in rep_data['section_n']['verified_items_summary'])
        assert rep_data['meta']['opinion_readiness'] == "DRAFT_PENDING_PROFESSIONAL_REVIEW"
        assert rep_data['meta']['is_final_certified'] is False

    # -------------------------------------------------------------------------
    # TEST 4: Vendor Control Account versus Subledger Variance
    # -------------------------------------------------------------------------
    def test_vendor_control_account_versus_subledger_variance(self, setup_engagement):
        engagement = setup_engagement['engagement']
        partner = setup_engagement['partner']

        # GL Trade Payables is ₹2,500,000.00.
        # Create subledger aging items totaling only ₹1,800,000.00 (Variance = ₹700,000 > perf mat ₹375,000).
        SubledgerAgingItem.objects.create(
            engagement=engagement,
            item_type='PAYABLE',
            party_name='Vendor ABC Logistics',
            invoice_number='INV-001',
            invoice_date=date(2026, 1, 15),
            due_date=date(2026, 2, 15),
            total_amount=Decimal('1800000.00'),
            outstanding_amount=Decimal('1800000.00'),
            ageing_bucket='31-60 days'
        )

        fin_readiness = AuditFinalizationService.evaluate_finalization_readiness(engagement, user=partner)
        assert fin_readiness['can_finalize'] is False
        assert any("Trade Payables subledger variance" in b for b in fin_readiness['blockers'])

    # -------------------------------------------------------------------------
    # TEST 5: Missing GST / Revenue Evidence Blocks Finalization
    # -------------------------------------------------------------------------
    def test_missing_gst_revenue_evidence_blocks_finalization(self, setup_engagement):
        engagement = setup_engagement['engagement']
        partner = setup_engagement['partner']

        # Ensure no GSTR-1 document exists for engagement
        engagement.documents.all().delete()

        fin_readiness = AuditFinalizationService.evaluate_finalization_readiness(engagement, user=partner)
        assert fin_readiness['can_finalize'] is False
        assert any("revenue verification evidence missing" in b.lower() for b in fin_readiness['blockers'])

        # Check ReportBuilder output
        rep_data = ConsolidatedAuditReportBuilder.build_report(engagement)
        assert rep_data['section_b']['revenue_verification_status'] == 'UNVERIFIED_MISSING_EVIDENCE'
        assert any("Revenue verification: UNVERIFIED" in item for item in rep_data['section_n']['verified_items_summary'])

    # -------------------------------------------------------------------------
    # TEST 6: Open High and Critical Exceptions Block Finalization
    # -------------------------------------------------------------------------
    def test_open_high_and_critical_exceptions_block_finalization(self, setup_engagement):
        engagement = setup_engagement['engagement']
        partner = setup_engagement['partner']

        # Create open critical exception
        exc = ExceptionItem.objects.create(
            engagement=engagement,
            finding_reference="EXC-CRIT-001",
            title="Unrecorded vendor liability under section 43B(h)",
            category="STATUTORY",
            severity=ExceptionSeverity.CRITICAL,
            status=ExceptionStatus.OPEN,
            amount_involved=Decimal("600000.00"),
            explanation="Unpaid MSME liability exceeding 45 days limit",
            recommended_action="Post adjusting entry and disallow under 43B(h)"
        )

        fin_readiness = AuditFinalizationService.evaluate_finalization_readiness(engagement, user=partner)
        assert fin_readiness['can_finalize'] is False
        assert any("critical audit exception(s) remain unresolved" in b for b in fin_readiness['blockers'])

        # Resolve the exception
        exc.status = ExceptionStatus.RESOLVED
        exc.resolution_notes = "Adjusting journal entry posted and liability acknowledged."
        exc.save()

        fin_readiness_resolved = AuditFinalizationService.evaluate_finalization_readiness(engagement, user=partner)
        assert not any("critical audit exception(s) remain unresolved" in b for b in fin_readiness_resolved['blockers'])

    # -------------------------------------------------------------------------
    # TEST 7: Workpaper Approval versus Finalization Eligibility
    # -------------------------------------------------------------------------
    def test_workpaper_approval_versus_finalization_eligibility(self, setup_engagement):
        engagement = setup_engagement['engagement']
        partner = setup_engagement['partner']

        # Add unapproved workpaper
        wp = Workpaper.objects.create(
            engagement=engagement,
            wp_reference="WP-REV-01",
            title="Substantive Revenue Testing",
            audit_area="REVENUE",
            status=WorkpaperStatus.PREPARED,
            preparer=partner
        )

        checklist = EngagementCompletionService.evaluate_checklist(engagement)
        chk06 = next(i for i in checklist['checklist_items'] if i['code'] == 'CHK-06-WORKPAPERS')
        assert chk06['is_satisfied'] is False

        # Approve workpaper
        wp.status = WorkpaperStatus.APPROVED
        wp.reviewed_by = partner
        wp.save()

        checklist_approved = EngagementCompletionService.evaluate_checklist(engagement)
        chk06_approved = next(i for i in checklist_approved['checklist_items'] if i['code'] == 'CHK-06-WORKPAPERS')
        assert chk06_approved['is_satisfied'] is True

    # -------------------------------------------------------------------------
    # TEST 8: Report Status Consistency Across PDF, XLSX, CSV, and API
    # -------------------------------------------------------------------------
    def test_report_status_consistency_across_pdf_xlsx_csv_api(self, setup_engagement):
        engagement = setup_engagement['engagement']
        engagement.is_finalized = False
        engagement.save()

        # 1. API representation
        rep_data = ConsolidatedAuditReportBuilder.build_report(engagement)
        assert rep_data['meta']['is_final_certified'] is False
        assert rep_data['meta']['opinion_readiness'] == "DRAFT_PENDING_PROFESSIONAL_REVIEW"
        assert rep_data['section_n']['statutory_status'] == "DRAFT_PENDING_PROFESSIONAL_REVIEW"

        # 2. PDF representation
        pdf_bytes = ReportExporterService.export_consolidated_pdf(engagement)
        assert isinstance(pdf_bytes, bytes)
        assert len(pdf_bytes) > 2000
        assert pdf_bytes.startswith(b"%PDF")

        # 3. Excel representation
        excel_bytes = ReportExporterService.export_consolidated_excel(engagement)
        assert isinstance(excel_bytes, bytes)
        assert excel_bytes.startswith(b"PK")

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        assert "Cover" in wb.sheetnames
        assert "Executive Summary" in wb.sheetnames
        assert "Audit Opinion & Governance" in wb.sheetnames

        ws_cover = wb["Cover"]
        assert "DRAFT AUDIT FILE" in str(ws_cover["B2"].value)

        ws_sum = wb["Executive Summary"]
        sum_rows = {row[0]: row[1] for row in ws_sum.iter_rows(min_row=4, values_only=True) if row[0]}
        assert "DRAFT" in str(sum_rows.get("Statutory Opinion Status", ""))
        assert "BLOCKED" in str(sum_rows.get("Report Integrity Verdict", ""))

        ws_op = wb["Audit Opinion & Governance"]
        assert "NO STATUTORY OPINION EXPRESSED" in str(ws_op["B4"].value)

        # 4. CSV representation
        csv_str = ReportExporterService.export_consolidated_csv(engagement)
        assert isinstance(csv_str, str)
        assert "DRAFT / UNFINALIZED — AUDIT REVIEW REQUIRED" in csv_str or "DRAFT / UNFINALIZED" in csv_str
        assert "DRAFT - PENDING INDEPENDENT REVIEW" in csv_str
        assert "BLOCKED" in csv_str
